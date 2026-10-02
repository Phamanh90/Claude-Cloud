"""
Test cho io-list-check: dựng 1 file Excel I/O list giả lập có cài sẵn lỗi đã biết,
chạy checker và khẳng định từng rule bắt đúng lỗi.

Chạy:  python -m pytest .claude/skills/io-list-check/tests -q
Test PDF thật (tuỳ chọn): đặt biến môi trường IO_SAMPLE_PDF=<đường dẫn PDF>
"""
import os
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import check  # noqa: E402

HEAD1 = ["NO.", "FIELD INSTRUMENT TAG NO.", "LOOP TAG NO. (DCS)", "SERVICE", "FROM/ TỪ", "TO/ TỚI", "P&ID NO.",
         "SIGNAL", None, None, None, None, "RANGE", None, None, "INPUT/OUTPUT DEVICE", None, None,
         "FUNCTION", None, None, None, None, None, None, None, "REMARKS"]
HEAD2 = [None] * 7 + ["TYPE", "SPEC.", "CONTACT", "ACTION", "CIRCUIT VOLTAGE", "LOWER", "UPPER", "UNIT",
                      "NODE NO.", "SLOT NO.", "CHANNEL NO.", "INTERLOCK", "ALARM", None, None, None, None,
                      "CONTROL ACTION", "BURNOUT SETTING", None]
HEAD3 = [None] * 19 + ["AL", "HH", "H", "L", "LL", None, None, None]


def row(no, ftag, ltag, typ, spec="---", contact="---", volt="---", lo="---", hi="---", unit="---",
        node="", slot="", ch="", bo="---"):
    r = [no, ftag, ltag, f"{ltag} service", "Field", "DCS", "PID-1", typ, spec, contact, "---", volt,
         lo, hi, unit, node, slot, ch, "", "", "", "", "", "", "", bo, ""]
    return r


def build(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ABBR"
    for t, d in [("AI", "Analog Input"), ("AO", "Analog Output"), ("DI", "Discrete Input"),
                 ("DO", "Discrete Output"), ("COM-DI", "Discrete Input")]:
        ws.append(["ABBREVIATIONS", f"{t} : {d}"])

    ws = wb.create_sheet("SUMMARY")
    ws.append(["SYSTEM / HỆ THỐNG", "AI", "AO", "DI", "DO", "PI"])
    ws.append(["DCS-Tanks", 4, 1, 3, 1, 0])     # thực tế DI = 4 -> lệch
    ws.append(["Total", 4, 1, 3, 1, 0])

    ws = wb.create_sheet("IO_TANK")
    ws.append([None, None, None, "SIGNAL I/O LIST (DCS TANK / HARD-WIRED - HỆ THỐNG BỒN)"])
    ws.append(HEAD1)
    ws.append(HEAD2)
    ws.append(HEAD3)
    rows = [
        row(1, "PT-101", "PIA-101", "AI", "4-20mA", lo="0,0", hi="1,0", unit="MPaG", node="1", slot="1", ch="1", bo="U"),
        row(2, "PT-102", "PIA-102", "AI", "4-20mA", lo="5", hi="1", unit="MPaG", node="1", slot="1", ch="2", bo="U"),  # R04 lower>=upper
        row(3, "TT-103", "TI-103", "AI", "4-20mA", lo="Hold", hi="Hold", unit="°C", node="1", slot="1", ch="2"),  # R06, R04 burnout, R07 trùng kênh
        row(4, "LS-101", "LAH-101", "DI", "N.C", "S", "D (DC24V)", node="1", slot="2", ch="1"),
        row(5, "LS-102", "LSH-102", "AI", "4-20mA", lo="0", hi="1", unit="m", node="1", slot="1", ch="3", bo="U"),  # R05 switch nhưng AI
        row(6, "", "ZSO-101", "DI", "N.O", "S", "D (DC24V)", node="1", slot="2", ch="2"),
        row(7, "", "ZSO-101", "DI", "N.O", "S", "D (DC24V)", node="1", slot="2", ch="3"),  # R01 trùng trong section
        row(8, "", "HS-101", "DI", "XX", "S", "D (DC24V)", node="1", slot="2", ch="4"),   # R04 spec sai
        row(9, "", "HSS-101", "DO", "N.O", "M", "D (DC24V)", node="1", slot="3", ch="1"),
        row(10, "", "PIC-101", "AO", "4-20mA", lo="0", hi="100", unit="%", node="1", slot="4", ch="1"),
        row(11, "", "XA-101", "MOD", "-", "-", "-"),  # R03 kiểu lạ
    ]
    for r in rows:
        ws.append(r)

    ws = wb.create_sheet("IO_TANK_COM")
    ws.append([None, None, None, "SIGNAL I/O LIST (DCS TANK / COMMUNICATION - GIAO TIẾP)"])
    ws.append(HEAD1)
    ws.append(HEAD2)
    ws.append(HEAD3)
    ws.append(row(1, "MV-101", "ZSO-101", "COM-DI", "N.O", "S"))  # trùng tag khác section
    wb.save(path)


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    p = tmp_path_factory.mktemp("io") / "sample.xlsx"
    build(p)
    cfg = check.load_config(tag_scope="section")
    io, findings, recount = check.run(str(p), cfg)
    return io, findings, recount, p


def has(findings, rule, sev=None, text=""):
    return any(f["rule"] == rule and (sev is None or f["severity"] == sev) and text in f["message"] for f in findings)


def test_extract_excel(result):
    io, _, _, _ = result
    assert len(io.records) == 12
    assert io.sections == ["DCS TANK / HARD-WIRED", "DCS TANK / COMMUNICATION"]
    assert {"AI", "DI", "COM-DI"} <= io.abbreviations
    assert io.records[0].get("lower") == "0,0" and io.records[0].get("burnout") == "U"


def test_r01_section_scope(result):
    _, f, _, _ = result
    assert has(f, "R01", "High", "trong cùng section")
    assert has(f, "R01", "Info", "2 section")


def test_r01_document_scope(result):
    _, _, _, p = result
    _, f, _ = check.run(str(p), check.load_config(tag_scope="document"))
    assert has(f, "R01", "High", "trùng 3 lần")


def test_r02_summary(result):
    _, f, recount, _ = result
    assert has(f, "R02", "High", "cột DI: Summary = 3, đếm chi tiết (DCS TANK / HARD-WIRED) = 4")
    assert all(r["ok"] for r in recount if r["type"] != "DI")


def test_r03_r04_r05_r06_r07(result):
    _, f, _, _ = result
    assert has(f, "R03", "Medium", "'MOD'")
    assert has(f, "R04", "Medium", "lower (5) ≥ upper (1)")
    assert has(f, "R04", "Medium", "burnout")
    assert has(f, "R04", "Medium", "spec phải là N.O/N.C (hiện tại 'XX')")
    assert has(f, "R05", "Medium", "Tag LSH")
    assert not has(f, "R05", text="Tag PIA") and not has(f, "R05", text="Tag HSS")
    assert has(f, "R06", "Info", "lower, upper")
    assert has(f, "R07", "High", "1/1/2")


def test_report_written(result, tmp_path):
    io, f, rc, _ = result
    out = tmp_path / "Claude_IOCheck_sample.xlsx"
    check.write_report(io, f, rc, out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Findings", "Summary_Recount", "Data"]
    assert wb["Data"].max_row == 13


@pytest.mark.skipif(not os.environ.get("IO_SAMPLE_PDF"), reason="cần IO_SAMPLE_PDF")
def test_pdf_sample():
    io, f, _ = check.run(os.environ["IO_SAMPLE_PDF"], check.load_config(tag_scope="section"))
    assert len(io.records) > 0
