"""
check.py — Kiểm tra I/O list (PDF/Excel) và xuất báo cáo Excel.

Cách dùng:
  python check.py <input.pdf|xlsx> [--out DIR] [--tag-scope document|section]
                  [--map "Tên trong Summary=Tên section"] [--config rules.json]

Kết quả: <DIR>/Claude_IOCheck_<tên file>.xlsx (sheet Findings / Summary_Recount / Data)
và tóm tắt in ra màn hình.
"""
import argparse
import collections
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract  # noqa: E402

SKILL_DIR = Path(__file__).resolve().parent.parent
SEV_ORDER = {"High": 0, "Medium": 1, "Low": 2, "Info": 3}


# ---------------------------------------------------------------------------
# Tiện ích
# ---------------------------------------------------------------------------
def norm(v):
    return (v or "").strip()


def is_empty(v, cfg):
    return norm(v).lower() in cfg["empty_values"]


def is_placeholder(v, cfg):
    return norm(v).lower() in cfg["placeholder_values"]


def to_num(v):
    s = norm(v).replace(" ", "")
    # Hỗ trợ dấu thập phân kiểu VN "0,0" và kiểu EN "0.0"
    if re.fullmatch(r"-?\d+(,\d+)?", s):
        s = s.replace(",", ".")
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def base_type(t):
    t = norm(t).upper()
    return t[4:] if t.startswith("COM-") else t


def is_comm(t):
    return norm(t).upper().startswith("COM-")


class Findings:
    def __init__(self):
        self.items = []
        self._groups = {}

    def group(self, sev, rule, key, rec, msg):
        """Gom các finding cùng bản chất (cùng key) thành 1 dòng kèm danh sách vị trí."""
        g = self._groups.get((rule, key))
        if g is None:
            g = self._groups[(rule, key)] = {"sev": sev, "rec": rec, "msg": msg, "locs": [], "tags": []}
        g["locs"].append(rec.loc)
        g["tags"].append(rec.get("loop_tag"))

    def flush(self, max_locs=12):
        for (rule, _), g in self._groups.items():
            n = len(g["locs"])
            locs = ", ".join(g["locs"][:max_locs]) + (f" … (+{n - max_locs})" if n > max_locs else "")
            tag = g["tags"][0] if n == 1 else f"{g['tags'][0]} (+{n - 1})"
            self.add(g["sev"], rule, g["rec"], f"{g['msg']} — {n} tín hiệu: {locs}", tag=tag)
        self._groups = {}

    def add(self, sev, rule, rec, msg, tag=None):
        self.items.append({
            "severity": sev, "rule": rule,
            "loc": rec.loc if rec else "", "section": rec.section if rec else "",
            "tag": tag if tag is not None else (rec.get("loop_tag") if rec else ""),
            "type": rec.get("type") if rec else "", "message": msg,
        })


# ---------------------------------------------------------------------------
# R01 — Trùng Loop Tag
# ---------------------------------------------------------------------------
def r01_duplicate_tags(io, cfg, f):
    by_tag = collections.defaultdict(list)
    for r in io.records:
        t = norm(r.get("loop_tag")).upper()
        if not is_empty(t, cfg):
            by_tag[t].append(r)
    for tag, recs in by_tag.items():
        if len(recs) < 2:
            continue
        sections = {r.section for r in recs}
        locs = ", ".join(r.loc for r in recs)
        if cfg["tag_scope"] == "section":
            # Trùng trong cùng section mới là lỗi; khác section -> chỉ ghi nhận
            per_sec = collections.Counter(r.section for r in recs)
            for sec, n in per_sec.items():
                if n > 1:
                    rs = [r for r in recs if r.section == sec]
                    f.add("High", "R01", rs[0], f"Loop tag trùng {n} lần trong cùng section: {', '.join(r.loc for r in rs)}")
            if len(sections) > 1:
                f.add("Info", "R01", recs[0], f"Loop tag dùng ở {len(sections)} section (hợp lệ theo quy ước dự án): {locs}")
        else:
            f.add("High", "R01", recs[0], f"Loop tag trùng {len(recs)} lần: {locs}")


# ---------------------------------------------------------------------------
# R02 — Đếm lại so với I/O SUMMARY
# ---------------------------------------------------------------------------
def _tokens(s):
    s = re.sub(r"[_/\-]", " ", s.upper())
    stop = {"DCS", "RIO", "PLC", "AREA", "HARD", "WIRED", "COMMUNICATION", "SYSTEM"}
    out = set()
    for w in s.split():
        if w in stop:
            continue
        if len(w) > 3 and w.endswith("S") and not w.isdigit():
            w = w[:-1]  # TANKS -> TANK
        out.add(w)
    return out


def map_summary(name, sections, overrides):
    """Map tên hệ trong bảng Summary -> section trong danh sách chi tiết."""
    if name in overrides:
        return overrides[name]
    st = _tokens(name)
    best, best_key = None, None
    for sec in sections:
        tt = _tokens(sec)
        nums_sec = {t for t in tt if any(ch.isdigit() for ch in t)}
        nums_sum = {t for t in st if any(ch.isdigit() for ch in t)}
        if nums_sec != nums_sum:
            continue  # số hiệu khu vực (VC-5002...) phải khớp tuyệt đối
        inter = len(st & tt)
        if inter == 0:
            continue
        key = (-inter, len(tt - st))
        if best_key is None or key < best_key:
            best, best_key = sec, key
    return best


def r02_summary(io, cfg, f, overrides):
    rows = []
    if not io.summary:
        f.add("Info", "R02", None, "Không tìm thấy bảng I/O SUMMARY — bỏ qua đối chiếu.", tag="")
        return rows
    hw_secs = [s for s in io.sections if "COMMUNICATION" not in s.upper()]
    count = collections.Counter((r.section, norm(r.get("type")).upper()) for r in io.records)

    for table in ("hardwired", "comm"):
        entries = [e for e in io.summary if e["table"] == table]
        if not entries:
            continue
        types = cfg["summary_types_hardwired"] if table == "hardwired" else cfg["summary_types_comm"]
        total_row = next((e for e in entries if e["name"].strip().upper() in ("TOTAL", "TỔNG", "TỔNG CỘNG")), None)
        body = [e for e in entries if e is not total_row]

        # Tự kiểm tra cộng hàng của chính bảng Summary
        if total_row:
            for t in types:
                s = sum(e["counts"].get(t, 0) for e in body)
                if s != total_row["counts"].get(t, 0):
                    f.add("Medium", "R02", None, f"Bảng Summary ({table}) cộng sai cột {t}: tổng các hàng = {s}, hàng Total = {total_row['counts'].get(t, 0)} ({total_row['loc']})", tag="SUMMARY")

        if table == "hardwired":
            for e in body:
                sec = map_summary(e["name"], hw_secs, overrides)
                for t in types:
                    want = e["counts"].get(t, 0)
                    got = count[(sec, t)] if sec else None
                    rows.append({"table": table, "summary_name": e["name"], "section": sec or "(không map được)",
                                 "type": t, "summary": want, "recount": got, "ok": got == want})
                    if sec and got != want:
                        f.add("High", "R02", None, f"'{e['name']}' cột {t}: Summary = {want}, đếm chi tiết ({sec}) = {got}", tag=e["name"])
                if not sec:
                    f.add("Medium", "R02", None, f"Không map được '{e['name']}' sang section nào — dùng --map để chỉ định", tag=e["name"])
        else:
            # Bảng COM thường chia theo giao thức/thiết bị, không trùng section -> đối chiếu tổng
            got_all = collections.Counter(norm(r.get("type")).upper() for r in io.records if is_comm(r.get("type")))
            if total_row:
                for t in types:
                    want, got = total_row["counts"].get(t, 0), got_all.get(t, 0)
                    rows.append({"table": table, "summary_name": "Total", "section": "(toàn bộ tín hiệu COM)",
                                 "type": t, "summary": want, "recount": got, "ok": got == want})
                    if got != want:
                        f.add("High", "R02", None, f"Tổng {t}: Summary = {want}, đếm chi tiết = {got}", tag="SUMMARY")
    return rows


# ---------------------------------------------------------------------------
# R03 — Kiểu tín hiệu / Spec / Contact có trong bảng viết tắt
# ---------------------------------------------------------------------------
def r03_abbreviations(io, cfg, f):
    allowed = io.abbreviations or set(cfg["signal_types_default"])
    for r in io.records:
        t = norm(r.get("type")).upper()
        if t not in allowed:
            f.group("Medium", "R03", (r.section, t), r,
                    f"Kiểu tín hiệu '{r.get('type') or '(trống)'}' không có trong bảng chữ viết tắt ({', '.join(sorted(allowed))})")


# ---------------------------------------------------------------------------
# R04 — Logic theo kiểu tín hiệu
# ---------------------------------------------------------------------------
def r04_type_logic(io, cfg, f):
    spec_ok = set(cfg["discrete_spec"])
    contact_ok = set(cfg["discrete_contact"])
    bo_ok = set(cfg["burnout_values"])
    for r in io.records:
        t = norm(r.get("type")).upper()
        bt, comm = base_type(t), is_comm(t)
        bo = norm(r.get("burnout")).upper()

        if bt in ("AI", "AO") and not comm:
            lo, hi, unit = r.get("lower"), r.get("upper"), r.get("unit")
            if is_placeholder(lo, cfg) or is_placeholder(hi, cfg):
                pass  # để R06 xử lý
            elif is_empty(lo, cfg) or is_empty(hi, cfg):
                f.add("Medium", "R04", r, f"{bt} thiếu dải đo (lower='{lo}', upper='{hi}')")
            else:
                a, b = to_num(lo), to_num(hi)
                if a is None or b is None:
                    f.add("Medium", "R04", r, f"Dải đo không phải số: lower='{lo}', upper='{hi}'")
                elif a >= b:
                    f.add("Medium", "R04", r, f"Dải đo sai: lower ({lo}) ≥ upper ({hi})")
            if is_empty(unit, cfg) and not is_placeholder(lo, cfg):
                f.add("Low", "R04", r, f"{bt} thiếu đơn vị (unit)")

        if bt == "AI" and not comm and bo not in bo_ok:
            f.add("Medium", "R04", r, f"AI chưa có burnout setting U/D (hiện tại '{r.get('burnout')}')")
        if bt in ("AO", "DI", "DO", "PI") and bo in bo_ok:
            f.add("Low", "R04", r, f"{t} không cần burnout setting nhưng đang ghi '{bo}'")

        if bt in ("DI", "DO"):
            spec = re.sub(r"[\s.]", "", norm(r.get("spec")).upper())
            if spec not in spec_ok:
                f.add("Medium", "R04", r, f"{t} spec phải là N.O/N.C (hiện tại '{r.get('spec')}')")
            if not comm:
                if norm(r.get("contact")).upper() not in contact_ok:
                    f.add("Low", "R04", r, f"{t} contact phải là M/S (hiện tại '{r.get('contact')}')")
                if is_empty(r.get("voltage"), cfg):
                    f.add("Low", "R04", r, f"{t} thiếu circuit voltage")


# ---------------------------------------------------------------------------
# R05 — Tag ↔ kiểu tín hiệu theo ISA-5.1
# ---------------------------------------------------------------------------
def isa_expected(prefix, cfg):
    """Từ chữ cái tag ISA-5.1 suy ra tập kiểu tín hiệu (base) hợp lệ.
    Trả về (set, giải thích) hoặc (None, lý do) nếu không áp dụng."""
    first, rest = prefix[0], prefix[1:]
    # Bỏ modifier của chữ cái đầu (D: differential, F: ratio, J: scan, K: time rate, Q: totalize)
    while rest and rest[0] in cfg["isa"]["first_letter_modifiers"]:
        rest = rest[1:]
    # PSV/PSE: S = safety (modifier), không phải switch
    if rest[:2] in ("SV", "SE"):
        return {"DI", "DO", "AI"}, "S = safety"
    if "S" in rest:
        # Switch: phần sau S là modifier trạng thái (O/C/H/L/R/P/A...) -> bỏ qua
        if first == "H":
            return {"DI", "DO"}, "H…S = hand switch (vào hoặc lệnh ra)"
        return {"DI"}, f"{first}S = switch → DI"
    if "C" in rest:
        return {"AI", "AO", "DO"}, "C = controller"
    if any(ch in rest for ch in "VYZ"):
        return {"AO", "DO"}, "V/Y/Z = final element/relay/actuator → output"
    if any(ch in rest for ch in "TEIR"):
        return {"AI", "PI"}, "T/E/I/R = transmit/element/indicate/record → analog in"
    if "A" in rest:
        return {"AI", "DI"}, "A = alarm"
    return None, "không suy được chức năng"


def r05_isa(io, cfg, f):
    rx = re.compile(cfg["isa"]["tag_regex"])
    non_isa = {p.upper() for p in cfg["isa"]["non_isa_prefixes"]}
    skipped = collections.Counter()
    for r in io.records:
        tag = norm(r.get("loop_tag")).upper()
        bt = base_type(r.get("type"))
        if bt not in ("AI", "AO", "DI", "DO", "PI"):
            continue  # kiểu tín hiệu lạ đã báo ở R03
        m = rx.match(tag)
        if not m or m.group(1) in non_isa:
            skipped[tag.split("-")[0] or "(trống)"] += 1
            continue
        exp, why = isa_expected(m.group(1), cfg)
        if exp is None:
            skipped[m.group(1)] += 1
            continue
        if bt not in exp:
            f.group("Medium", "R05", (m.group(1), norm(r.get("type")).upper()), r,
                    f"Tag {m.group(1)} ({why}) nhưng kiểu tín hiệu là {r.get('type')} — kiểm tra lại tag hoặc type")
    if skipped:
        f.add("Info", "R05", None, "Không áp ISA-5.1 cho tag ngoài chuẩn: " + ", ".join(f"{k}×{v}" for k, v in skipped.most_common()), tag="")


# ---------------------------------------------------------------------------
# R06 — Ô còn Hold/TBC
# ---------------------------------------------------------------------------
def r06_placeholders(io, cfg, f):
    for r in io.records:
        cols = [k for k, v in r.data.items() if is_placeholder(v, cfg)]
        if cols:
            f.group("Info", "R06", (r.section, tuple(cols)), r, f"Còn giá trị Hold/TBC ở cột: {', '.join(cols)}")


# ---------------------------------------------------------------------------
# R07 — Gán kênh I/O (chỉ khi đã có Node/Slot/Channel)
# ---------------------------------------------------------------------------
def r07_channels(io, cfg, f):
    assigned = [r for r in io.records if not is_empty(r.get("channel"), cfg)]
    if not assigned:
        f.add("Info", "R07", None, f"Chưa gán Node/Slot/Channel cho {len(io.records)} tín hiệu — bỏ qua kiểm tra kênh", tag="")
        return
    used = collections.defaultdict(list)
    for r in assigned:
        key = (norm(r.get("node")), norm(r.get("slot")), norm(r.get("channel")))
        used[key].append(r)
    for key, rs in used.items():
        if len(rs) > 1:
            f.add("High", "R07", rs[0], f"Trùng kênh Node/Slot/Ch {'/'.join(key)}: {', '.join(r.loc for r in rs)}")
    # Một card (node, slot) chỉ nên chứa một kiểu tín hiệu
    card_types = collections.defaultdict(set)
    for r in assigned:
        card_types[(norm(r.get("node")), norm(r.get("slot")))].add(base_type(r.get("type")))
    for key, ts in card_types.items():
        if len(ts) > 1:
            f.add("Medium", "R07", None, f"Card Node/Slot {'/'.join(key)} chứa nhiều kiểu tín hiệu: {', '.join(sorted(ts))}", tag="")


# ---------------------------------------------------------------------------
# Xuất báo cáo
# ---------------------------------------------------------------------------
def write_report(io, findings, recount, out_path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9D9D9")

    def header(ws, cols, widths):
        ws.append(cols)
        for c, w in zip(ws[1], widths):
            c.font, c.fill = bold, head_fill
            ws.column_dimensions[c.column_letter].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Findings"
    header(ws, ["Severity", "Rule", "Location", "Section", "Loop Tag", "Type", "Message"],
           [10, 7, 16, 34, 18, 9, 90])
    for it in sorted(findings, key=lambda x: (SEV_ORDER[x["severity"]], x["rule"])):
        ws.append([it["severity"], it["rule"], it["loc"], it["section"], it["tag"], it["type"], it["message"]])
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Summary_Recount")
    header(ws, ["Table", "Summary name", "Mapped section", "Type", "Summary", "Recount", "OK"],
           [10, 28, 40, 8, 9, 9, 6])
    for r in recount:
        ws.append([r["table"], r["summary_name"], r["section"], r["type"], r["summary"], r["recount"], "OK" if r["ok"] else "LỆCH"])

    ws = wb.create_sheet("Data")
    keys = [k for k, _, _ in extract.COLUMN_KEYS]
    header(ws, ["Location", "Section"] + keys, [14, 34] + [12] * len(keys))
    for r in io.records:
        ws.append([r.loc, r.section] + [r.get(k) for k in keys])

    wb.save(out_path)


RULES = [r01_duplicate_tags, r03_abbreviations, r04_type_logic, r05_isa, r06_placeholders, r07_channels]


def run(path, cfg, overrides=None):
    io = extract.load(path, cfg)
    f = Findings()
    recount = r02_summary(io, cfg, f, overrides or {})
    for rule in RULES:
        rule(io, cfg, f)
    f.flush()
    return io, f.items, recount


def load_config(path=None, **override):
    cfg = json.loads((Path(path) if path else SKILL_DIR / "references" / "rules.json").read_text(encoding="utf-8"))
    cfg.update({k: v for k, v in override.items() if v is not None})
    return cfg


def main(argv=None):
    ap = argparse.ArgumentParser(description="Kiểm tra I/O list (PDF/Excel)")
    ap.add_argument("input")
    ap.add_argument("--out", default=".")
    ap.add_argument("--tag-scope", choices=["document", "section"], help="Phạm vi duy nhất của loop tag")
    ap.add_argument("--map", action="append", default=[], help='"Tên Summary=Tên section" (lặp lại được)')
    ap.add_argument("--config", help="File rules.json thay thế")
    a = ap.parse_args(argv)

    cfg = load_config(a.config, tag_scope=a.tag_scope)
    overrides = dict(m.split("=", 1) for m in a.map)
    io, findings, recount = run(a.input, cfg, overrides)

    out = Path(a.out) / f"Claude_IOCheck_{Path(a.input).stem}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    write_report(io, findings, recount, out)

    # Tóm tắt ngắn cho Claude/người dùng đọc (không in toàn bộ findings để tiết kiệm token)
    print(f"Nguồn: {io.source}")
    print(f"Tín hiệu: {len(io.records)} | Section: {len(io.sections)} | Hàng Summary: {len(io.summary)} | tag_scope={cfg['tag_scope']}")
    print("Kiểu tín hiệu khai báo:", ", ".join(sorted(io.abbreviations)) or "(không thấy — dùng mặc định)")
    sev = collections.Counter(x["severity"] for x in findings)
    by_rule = collections.Counter((x["rule"], x["severity"]) for x in findings)
    print("Findings:", ", ".join(f"{k}={sev[k]}" for k in SEV_ORDER if sev[k]))
    for (rule, s), n in sorted(by_rule.items(), key=lambda x: (x[0][0], SEV_ORDER[x[0][1]])):
        print(f"  {rule} {s}: {n}")
    print("Top High/Medium:")
    top = [x for x in sorted(findings, key=lambda x: (SEV_ORDER[x["severity"]], x["rule"])) if x["severity"] in ("High", "Medium")]
    for x in top[:15]:
        print(f"  [{x['severity']}] {x['rule']} {x['loc']} {x['tag']}: {x['message']}")
    if len(top) > 15:
        print(f"  ... còn {len(top) - 15} mục, xem file báo cáo")
    print(f"Báo cáo: {out}")


if __name__ == "__main__":
    main()
